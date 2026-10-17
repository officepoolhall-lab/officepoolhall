#!/usr/bin/env python3
"""Parse Brooks's NCAA Supercontest workbook into JSON for build.mjs.

    python -I parse.py <workbook.xlsx> <out.json> [--weeks 1-6] [--partial 7]

--partial N: week N (the last requested week) is in progress. Its games may have neither an
outcome nor a cover yet ("not decided"); every other check stays strict.

The workbook holds real people's names and emails: write <out.json> outside the repo.
Exits non-zero, listing every problem, if anything doesn't parse or validate.
Needs openpyxl (use a throwaway venv; it is not a project dependency).
"""
import argparse
import difflib
import json
import re
import sys

import openpyxl

RANK_RE = re.compile(r"^#(\d+)\s+")
SIDE_RE = re.compile(r"^(?P<name>.+?)(?:\s+(?P<spread>[+-]\d+(?:\.\d+)?))?$")
OUTCOME_RE = re.compile(r"^(?P<name>.+?)\s+(?P<w>\d+)\s*-\s*(?P<l>\d+)$")
RECORD_RE = re.compile(r"^(\d+)\s*-\s*(\d+)$")

problems = []
warnings = []


def ws_text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def parse_side(text):
    """'#14 Tennessee -4.5' -> {name, rank, spread}."""
    t = ws_text(text)
    rank = None
    m = RANK_RE.match(t)
    if m:
        rank = int(m.group(1))
        t = t[m.end():]
    m = SIDE_RE.match(t)
    if not m:
        return None
    spread = m.group("spread")
    return {"name": m.group("name").strip(), "rank": rank, "spread": float(spread) if spread else None}


def key(name):
    return re.sub(r"[^a-z0-9]", "", name.lower())


def find_team(teams, name):
    """Index of the team a (possibly shortened or misspelled) name refers to, or None."""
    k = key(name)
    exact = [i for i, t in enumerate(teams) if key(t["name"]) == k]
    if len(exact) == 1:
        return exact[0]
    prefix = [i for i, t in enumerate(teams) if key(t["name"]).startswith(k) or k.startswith(key(t["name"]))]
    if len(prefix) == 1:
        return prefix[0]
    scores = [difflib.SequenceMatcher(None, k, key(t["name"])).ratio() for t in teams]
    best = max(range(len(teams)), key=lambda i: scores[i])
    if scores[best] >= 0.8 and scores[1 - best] < 0.8:
        return best
    return None


def parse_record(v, where):
    m = RECORD_RE.match(ws_text(v))
    if not m:
        problems.append(f"{where}: record {v!r} is not W-L")
        return {"w": 0, "l": 0}
    return {"w": int(m.group(1)), "l": int(m.group(2))}


def week_range(spec):
    out = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out.update(range(int(a), int(b or a) + 1))
    return sorted(out)


def parse_roster(wb):
    ws = wb["Roster"]
    header = [ws_text(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    try:
        ci = {h: header.index(h) for h in ("Name", "Email", "Payment Received")}
    except ValueError as e:
        problems.append(f"Roster: missing header ({e})")
        return []
    roster, seen_e, seen_n = [], set(), set()
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = ws_text(r[ci["Name"]])
        if not name:
            continue
        email = ws_text(r[ci["Email"]]).lower()
        if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            problems.append(f"Roster row {i}: bad email")
        if email in seen_e:
            problems.append(f"Roster row {i}: duplicate email")
        if name in seen_n:
            problems.append(f"Roster row {i}: duplicate name {name!r}")
        seen_e.add(email)
        seen_n.add(name)
        roster.append({"name": name, "email": email, "paid": ws_text(r[ci["Payment Received"]]).lower() == "yes"})
    return roster


def parse_standings(wb):
    ws = wb["Standings"]
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        name = ws_text(r[1])
        if not name or name == "Entry":
            continue
        out.append({
            "name": name,
            "record": parse_record(r[2], f"Standings row {i}"),
            "segments": [parse_record(r[c], f"Standings row {i} segment {c - 3}") for c in (4, 5, 6)],
        })
    return out


def parse_week(wb, n, partial=False):
    ws = wb[f"Week {n}"]
    rows = list(ws.iter_rows(values_only=True))
    if ws_text(rows[1][0]) != "Matchup / Spread" or ws_text(rows[2][0]) != "Outcome" or ws_text(rows[4][0]) != "Entry":
        problems.append(f"Week {n}: unexpected sheet layout")
        return None
    games = []
    for col, cell in enumerate(rows[1]):
        if col < 2 or not ws_text(cell):
            continue
        where = f"Week {n} col {col}"
        raw = ws_text(cell)
        parts = re.split(r"\s+vs\.?\s+", raw)
        sides = [parse_side(p) for p in parts] if len(parts) == 2 else None
        if not sides or None in sides:
            problems.append(f"{where}: can't parse matchup {raw!r}")
            continue
        with_spread = [s for s in sides if s["spread"] is not None]
        if len(with_spread) != 1:
            problems.append(f"{where}: spread must be on exactly one side: {raw!r}")
            continue
        if with_spread[0]["spread"] % 1 == 0:
            warnings.append(f"{where}: whole-number spread {raw!r}")
        g = {"col": col, "raw": raw, "teams": sides, "outcome": None, "cover": None, "decided": True}

        def team_spread(i):
            s = sides[i]["spread"]
            return s if s is not None else -sides[1 - i]["spread"]

        out_raw = ws_text(rows[2][col])
        if partial and not out_raw and not ws_text(rows[3][col]):
            # In-progress week: not played (or not entered) yet. build.mjs imports ESPN's status.
            g["decided"] = False
            games.append(g)
            continue
        m = OUTCOME_RE.match(out_raw)
        if not m:
            problems.append(f"{where}: no/bad outcome {out_raw!r} for {raw!r}")
        else:
            wi = find_team(sides, RANK_RE.sub("", m.group("name")))
            if wi is None:
                problems.append(f"{where}: outcome winner {m.group('name')!r} is neither team of {raw!r}")
            else:
                ws_, ls_ = int(m.group("w")), int(m.group("l"))
                if ws_ < ls_:
                    problems.append(f"{where}: outcome {out_raw!r} winner score lower than loser")
                g["outcome"] = {"winner": sides[wi]["name"], "w_score": ws_, "l_score": ls_}
        cov = parse_side(rows[3][col]) if ws_text(rows[3][col]) else None
        if not cov:
            problems.append(f"{where}: no cover cell for {raw!r}")
        else:
            ci = find_team(sides, cov["name"])
            if ci is None or key(sides[ci]["name"]) != key(cov["name"]):
                problems.append(f"{where}: cover {rows[3][col]!r} is neither team of {raw!r}")
            elif cov["spread"] != team_spread(ci):
                problems.append(f"{where}: cover spread {rows[3][col]!r} doesn't match line {raw!r}")
            else:
                g["cover"] = {"team": sides[ci]["name"], "spread": cov["spread"]}
                if g["outcome"]:
                    # Recompute the cover from the score as a sanity check.
                    wi = [s["name"] for s in sides].index(g["outcome"]["winner"])
                    margin_w = g["outcome"]["w_score"] - g["outcome"]["l_score"]
                    adj = margin_w + team_spread(wi)  # winner's margin against the line
                    if adj == 0:
                        problems.append(f"{where}: push on {raw!r}")
                    elif (adj > 0) != (ci == wi):
                        problems.append(f"{where}: cover cell disagrees with score {out_raw!r} and line {raw!r}")
        games.append(g)
    by_col = {g["col"]: g for g in games}

    entries = []
    for i, r in enumerate(rows[5:], start=6):
        name = ws_text(r[0])
        if not name:
            continue
        where = f"Week {n} row {i} ({name})"
        picks = []
        for col, cell in enumerate(r):
            if col < 2 or not ws_text(cell):
                continue
            g = by_col.get(col)
            if not g:
                problems.append(f"{where}: pick {cell!r} in a column with no game")
                continue
            p = parse_side(cell)
            ti = find_team(g["teams"], p["name"]) if p else None
            if ti is None or key(g["teams"][ti]["name"]) != key(p["name"]):
                problems.append(f"{where}: pick {cell!r} isn't a team in {g['raw']!r}")
                continue
            s = g["teams"][ti]["spread"]
            exp = s if s is not None else -g["teams"][1 - ti]["spread"]
            if p["spread"] is not None and p["spread"] != exp:
                problems.append(f"{where}: pick {cell!r} spread doesn't match {g['raw']!r}")
            picks.append({"col": col, "raw": ws_text(cell), "team": g["teams"][ti]["name"], "spread": exp})
        rec = parse_record(r[1], where)
        # Sheet record vs picks graded by the sheet's own cover row.
        w = sum(1 for p in picks if by_col[p["col"]]["cover"] and by_col[p["col"]]["cover"]["team"] == p["team"])
        decided = sum(1 for p in picks if by_col[p["col"]]["cover"])
        if (w, decided - w) != (rec["w"], rec["l"]):
            problems.append(f"{where}: sheet record {rec['w']}-{rec['l']} but picks grade to {w}-{decided - w}")
        entries.append({"name": name, "record": rec, "picks": picks})

    # Footer rows: per-game count of winning picks, then losing picks.
    totals = [r for r in rows[5:] if not ws_text(r[0]) and any(isinstance(x, (int, float)) for x in r[2:])]
    if len(totals) >= 2:
        for g in games:
            if not g["cover"]:
                continue
            on = [p["team"] for e in entries for p in e["picks"] if p["col"] == g["col"]]
            won = sum(1 for t in on if t == g["cover"]["team"])
            c = [won, len(on) - won]
            sheet = [totals[0][g["col"]], totals[1][g["col"]]]
            if c != sheet:
                warnings.append(f"Week {n} col {g['col']}: pick counts {c} vs sheet footer {sheet}")
    return {"week": n, "games": games, "entries": entries}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--weeks", default="1-6")
    ap.add_argument("--partial", type=int, help="the last requested week, still in progress")
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    weeks = week_range(a.weeks)
    if a.partial is not None and a.partial != weeks[-1]:
        print(f"--partial {a.partial} must be the last requested week ({weeks[-1]})", file=sys.stderr)
        sys.exit(2)
    roster = parse_roster(wb)
    standings = parse_standings(wb)
    parsed = []
    for n in weeks:
        if f"Week {n}" not in wb.sheetnames:
            problems.append(f"Week {n}: no sheet")
            continue
        w = parse_week(wb, n, partial=n == a.partial)
        if w:
            w["partial"] = n == a.partial
            if not w["games"]:
                problems.append(f"Week {n}: no games")
            parsed.append(w)

    # Records of sheet weeks that weren't requested (so verify.mjs can reconcile with Standings).
    excluded = []
    for sn in wb.sheetnames:
        m = re.match(r"^Week (\d+)$", sn)
        if m and int(m.group(1)) not in weeks:
            ws = wb[sn]
            recs = {}
            for r in list(ws.iter_rows(values_only=True))[5:]:
                if ws_text(r[0]) and RECORD_RE.match(ws_text(r[1])):
                    recs[ws_text(r[0])] = parse_record(r[1], sn)
            excluded.append({"week": int(m.group(1)), "records": recs})

    names = {r["name"] for r in roster}
    for w in parsed:
        for e in w["entries"]:
            if e["name"] not in names:
                problems.append(f"Week {w['week']}: entry {e['name']!r} not on the Roster")
            if len(e["picks"]) not in (0, 5):
                warnings.append(f"Week {w['week']}: {e['name']} has {len(e['picks'])} picks")
    for s in standings:
        if s["name"] not in names:
            problems.append(f"Standings: {s['name']!r} not on the Roster")

    for wmsg in warnings:
        print("warning:", wmsg, file=sys.stderr)
    if problems:
        print(f"{len(problems)} problem(s):", file=sys.stderr)
        for p in problems:
            print("  -", p, file=sys.stderr)
        sys.exit(1)
    with open(a.out, "w") as f:
        json.dump({"roster": roster, "standings": standings, "weeks": parsed, "excluded_weeks": excluded,
                   "partial_week": a.partial}, f, indent=1)
    print(f"ok: {len(roster)} roster, {len(standings)} standings, weeks {weeks}: "
          + ", ".join(f"W{w['week']} {len(w['games'])} games/{sum(len(e['picks']) for e in w['entries'])} picks"
                      + (f" ({sum(g['decided'] for g in w['games'])} decided, in progress)" if w["partial"] else "")
                      for w in parsed))


if __name__ == "__main__":
    main()
